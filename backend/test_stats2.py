from database import SessionLocal
from main import generate_excel_file
import openpyxl

db = SessionLocal()
wb_stream = generate_excel_file(db)
wb_stream.seek(0)
wb = openpyxl.load_workbook(wb_stream)
ws = wb.active

for row in range(4, 5):
    sira = ws.cell(row=row, column=1).value
    son = ws.cell(row=row, column=13).value
    budama_yapilacak = ws.cell(row=row, column=15).value
    budama_ihale = ws.cell(row=row, column=16).value
    koridor_yapilacak = ws.cell(row=row, column=24).value
    koridor_ihale = ws.cell(row=row, column=25).value
    print(f"Row {row} | Sira: {sira} | Son Durum: {son} | Budama: {budama_yapilacak} | Budama Ihale: {budama_ihale}")

db.close()

import openpyxl

wb = openpyxl.load_workbook("test_out.xlsx")
ws = wb.active

for row in range(4, 10):
    sira = ws.cell(row=row, column=1).value
    son_durum = ws.cell(row=row, column=13).value
    agac_yapildi = ws.cell(row=row, column=14).value
    agac_yapilacak = ws.cell(row=row, column=15).value
    print(f"Row {row} | Sira: {sira} | Durum: {son_durum} | Yapildi: {agac_yapildi} | Yapilacak: {agac_yapilacak}")

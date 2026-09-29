from database import SessionLocal
from main import generate_excel_file
import openpyxl

db = SessionLocal()
wb_stream = generate_excel_file(db)
wb_stream.seek(0)
wb = openpyxl.load_workbook(wb_stream)
ws = wb.active

for row in range(4, 10):
    sira = ws.cell(row=row, column=1).value
    son = ws.cell(row=row, column=13).value
    budama_yapilacak = ws.cell(row=row, column=15).value
    koridor_yapilacak = ws.cell(row=row, column=24).value
    print(f"Row {row} | Sira: {sira} | Son Durum: {son} | Budama: {budama_yapilacak} | Koridor: {koridor_yapilacak}")

db.close()

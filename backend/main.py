from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from datetime import timedelta
import models, schemas, auth
import math
from io import BytesIO
from fastapi.responses import StreamingResponse
import openpyxl
import smtplib
from email.message import EmailMessage
import os
import pydantic
from database import engine, get_db

models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="Orman Yangınları API")

import subprocess

APP_VERSION = "1.0.1"
try:
    commit_hash = subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], text=True).strip()
    iso_date = subprocess.check_output(["git", "log", "-1", "--format=%cI"], text=True).strip()
    
    # Simple manual parse for Turkish date
    from datetime import datetime
    d = datetime.fromisoformat(iso_date)
    months = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
    date_str = f"{d.day} {months[d.month - 1]} {d.year}"
    
    APP_VERSION = f"{APP_VERSION} (Commit: {commit_hash}) - {date_str}"
except Exception:
    pass

@app.get("/api/version")
def get_version():
    return {"version": APP_VERSION}

# Setup CORS for React frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], # For dev only, use explicit frontend URL in prod
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.post("/api/login", response_model=auth.Token)
def login_for_access_token(user_credentials: schemas.UserLogin, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.username == user_credentials.username).first()
    if not user or not auth.verify_password(user_credentials.password, user.password_hash):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect username or password",
            headers={"WWW-Authenticate": "Bearer"},
        )
    access_token_expires = timedelta(minutes=auth.ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = auth.create_access_token(
        data={"sub": user.username, "role": user.role}, expires_delta=access_token_expires
    )
    return {"access_token": access_token, "token_type": "bearer", "username": user.username, "role": user.role}

# To be implemented: dependency to get current user from token
from fastapi.security import OAuth2PasswordBearer
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="api/login")

def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    from jose import JWTError, jwt
    try:
        payload = jwt.decode(token, auth.SECRET_KEY, algorithms=[auth.ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise HTTPException(status_code=401, detail="Invalid auth credentials")
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid auth credentials")
    user = db.query(models.User).filter(models.User.username == username).first()
    if user is None:
        raise HTTPException(status_code=401, detail="User not found")
    return user

import functools

@functools.lru_cache(maxsize=1)
def get_ihale_excel_map():
    import pandas as pd
    import os
    excel_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'alikemal_orman.xlsx')
    try:
        df = pd.read_excel(excel_path, sheet_name='ORMANLIK ALANDAN GEÇEN HATLAR', header=2)
        df['KORİDOR AÇMA '] = df['KORİDOR AÇMA '].fillna('')
        df['AĞAÇ BUDAMA'] = df['AĞAÇ BUDAMA'].fillna('')
        excel_map = {}
        for _, row in df.iterrows():
            hat = str(row.get('Hat İsmi', '')).strip().upper()
            om = str(row.get('Operasyon Merkezi', '')).strip().upper()
            if not hat: continue
            
            koridor = str(row.get('KORİDOR AÇMA ', '')).strip().upper()
            budama = str(row.get('AĞAÇ BUDAMA', '')).strip().upper()
            excel_map[(om, hat)] = {
                "koridor": koridor,
                "budama": budama
            }
        return excel_map
    except Exception as e:
        print("Excel okuma hatasi:", e)
        return {}

@functools.lru_cache(maxsize=1)
def get_toroslar_excel_map():
    import pandas as pd
    import os
    excel_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'Kopya Ormanlık Alan Bakım Takip - Toroslar.xlsx')
    try:
        df = pd.read_excel(excel_path, header=2)
        excel_map = {}
        for _, row in df.iterrows():
            hat = str(row.get('Hat İsmi', '')).strip().upper()
            om = str(row.get('Operasyon Merkezi', '')).strip().upper()
            if not hat or hat == 'NAN': continue
            
            def safe_str(val):
                if pd.isna(val): return ""
                if isinstance(val, float) and val.is_integer(): return str(int(val))
                return str(val)
                
            excel_map[(om, hat)] = {
                "sira_no": safe_str(row.get('Sıra No')),
                "dagitim_sirketi": safe_str(row.get('Dağıtım Şirketi')),
                "il": safe_str(row.get('İl')),
                "ilce": safe_str(row.get('İlçe')),
                "gerilim_seviyesi": safe_str(row.get('Gerilim Seviyesi\n(AG/OG)')),
                "hat_uzunlugu": safe_str(row.get('Hat Uzunluğu\n(Km)')),
                "mevcut_risk": safe_str(row.get('Mevcut Risk\n(1./2./3. Derece Riskli)')),
                "planlanan_bakim": safe_str(row.get('Planlanan Bakım Tarihi')).split()[0] if safe_str(row.get('Planlanan Bakım Tarihi')) else "",
                "gerceklesen_bakim": safe_str(row.get('Gerçekleşen Bakım Tarihi')).split()[0] if safe_str(row.get('Gerçekleşen Bakım Tarihi')) else "",
                "siparis_no": safe_str(row.get('SİPARİŞ NUMARASI')),
            }
        return excel_map
    except Exception as e:
        print("Toroslar Excel okuma hatasi:", e)
        return {}

@app.get("/api/lines", response_model=list[schemas.LineResponse])
def get_lines(il: str = None, oms: str = None, db: Session = Depends(get_db)):
    # This will be optimized to join interventions, but for now we fetch lines
    query = db.query(models.Line)
    if il and il != "Tümü":
        query = query.filter(models.Line.il == il)
    if oms and oms != "Tümü":
        query = query.filter(models.Line.operasyon_merkezi == oms)
    
    lines = query.all()
    
    # Fetch all interventions to aggregate counts based on LATEST status per asset_id
    interventions = db.query(
        models.Intervention.sira_no,
        models.Intervention.category,
        models.Intervention.asset_id,
        models.Intervention.status,
        models.Intervention.created_at
    ).order_by(models.Intervention.created_at.asc()).all()
    
    # Keep only the latest status for each (sira_no, category, asset_id)
    latest_interventions = {}
    for inv in interventions:
        key = (inv.sira_no, inv.category, inv.asset_id)
        latest_interventions[key] = inv.status
        
    stats_map = {}
    for (sira_no, cat, asset_id), stat in latest_interventions.items():
        if sira_no not in stats_map:
            stats_map[sira_no] = {}
        if cat not in stats_map[sira_no]:
            stats_map[sira_no][cat] = {'Yapıldı': 0, 'Yapılmadı': 0, 'Bekliyor': 0}
        
        if stat in stats_map[sira_no][cat]:
            stats_map[sira_no][cat][stat] += 1
        else:
            stats_map[sira_no][cat][stat] = 1
        
    def get_count(sira, cat, stat):
        return stats_map.get(sira, {}).get(cat, {}).get(stat, 0)

    from sqlalchemy.orm import class_mapper
    result = []
    for line in lines:
        l_dict = {c.key: getattr(line, c.key) for c in class_mapper(line.__class__).column_attrs}
        
        # Add aggregated fields
        l_dict['agac_budama_ok'] = get_count(line.sira_no, 'Ağaç Budama', 'Yapıldı')
        l_dict['agac_budama_nok'] = get_count(line.sira_no, 'Ağaç Budama', 'Yapılmadı') + get_count(line.sira_no, 'Ağaç Budama', 'Bekliyor')
        
        l_dict['guzergah_degisimi_ok'] = get_count(line.sira_no, 'Güzergah Değişimi', 'Yapıldı')
        l_dict['guzergah_degisimi_nok'] = get_count(line.sira_no, 'Güzergah Değişimi', 'Yapılmadı') + get_count(line.sira_no, 'Güzergah Değişimi', 'Bekliyor')
        
        l_dict['beton_dokumu_ok'] = get_count(line.sira_no, 'Beton Dökümü', 'Yapıldı')
        l_dict['beton_dokumu_nok'] = get_count(line.sira_no, 'Beton Dökümü', 'Yapılmadı') + get_count(line.sira_no, 'Beton Dökümü', 'Bekliyor')
        
        l_dict['koridor_acma_ok'] = get_count(line.sira_no, 'Koridor Açma', 'Yapıldı')
        l_dict['koridor_acma_nok'] = get_count(line.sira_no, 'Koridor Açma', 'Yapılmadı') + get_count(line.sira_no, 'Koridor Açma', 'Bekliyor')
        
        l_dict['operasyon_mudahalesi_ok'] = get_count(line.sira_no, 'Operasyon Müdahalesi', 'Yapıldı')
        l_dict['operasyon_mudahalesi_nok'] = get_count(line.sira_no, 'Operasyon Müdahalesi', 'Yapılmadı') + get_count(line.sira_no, 'Operasyon Müdahalesi', 'Bekliyor')
        
        # Add ihale var/yok status
        excel_map = get_ihale_excel_map()
        line_om = str(line.operasyon_merkezi).strip().upper()
        line_hat = str(line.hat_ismi).strip().upper()
        
        koridor_excel = excel_map.get((line_om, line_hat), {}).get("koridor", "YOK")
        budama_excel = excel_map.get((line_om, line_hat), {}).get("budama", "YOK")
        
        l_dict['ihale_koridor'] = koridor_excel != "YOK" and koridor_excel != ""
        l_dict['ihale_budama'] = budama_excel != "YOK" and budama_excel != ""
        
        result.append(l_dict)
        
    return result

@app.put("/api/interventions/bulk-status")
def update_interventions_bulk(payload: schemas.InterventionBulkUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role == "izleyici":
        raise HTTPException(status_code=403, detail="Forbidden")
    db.query(models.Intervention).filter(models.Intervention.id.in_(payload.ids)).update({models.Intervention.status: payload.status}, synchronize_session=False)
    db.commit()
    return {"message": "Success"}

@app.post("/api/interventions/bulk-delete")
def delete_interventions_bulk(payload: schemas.InterventionBulkDelete, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role == "izleyici":
        raise HTTPException(status_code=403, detail="Forbidden")
    # Normal users can only delete their own records
    query = db.query(models.Intervention).filter(models.Intervention.id.in_(payload.ids))
    if current_user.role != "admin":
        query = query.filter(models.Intervention.created_by == current_user.username)
        
    query.delete(synchronize_session=False)
    db.commit()
    return {"message": "Success"}


@app.get("/api/interventions/{sira_no}", response_model=list[schemas.InterventionResponse])
def get_interventions_for_line(sira_no: int, db: Session = Depends(get_db)):
    interventions = db.query(models.Intervention).filter(models.Intervention.sira_no == sira_no).order_by(models.Intervention.created_at.desc()).all()
    return interventions

@app.post("/api/interventions/{sira_no}", response_model=schemas.InterventionResponse)
def create_intervention(sira_no: int, intervention: schemas.InterventionCreate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role == "izleyici":
        raise HTTPException(status_code=403, detail="Forbidden")
    db_intervention = models.Intervention(
        **intervention.dict(),
        sira_no=sira_no,
        created_by=current_user.username
    )
    db.add(db_intervention)
    db.commit()
    db.refresh(db_intervention)
    return db_intervention

@app.get("/api/me")
def read_users_me(current_user: models.User = Depends(get_current_user)):
    import json
    prefs = current_user.dashboard_preferences
    if not prefs: prefs = "[]"
    try:
        prefs_obj = json.loads(prefs)
    except:
        prefs_obj = []
        
    return {
        "username": current_user.username, 
        "role": current_user.role, 
        "default_il": current_user.default_il, 
        "default_oms": current_user.default_oms,
        "dashboard_preferences": prefs_obj
    }

@app.put("/api/me/defaults")
def update_user_defaults(payload: schemas.UserDefaultsUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    current_user.default_il = payload.default_il
    current_user.default_oms = payload.default_oms
    db.commit()
    return {"message": "Defaults updated successfully"}

@app.delete("/api/interventions/{id}")
def delete_intervention(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role == "izleyici":
        raise HTTPException(status_code=403, detail="Forbidden")
    intervention = db.query(models.Intervention).filter(models.Intervention.id == id).first()
    if not intervention:
        raise HTTPException(status_code=404, detail="Not found")
    
    if current_user.role != "admin" and intervention.created_by != current_user.username:
        raise HTTPException(status_code=403, detail="Forbidden: You can only delete your own records")
        
    db.delete(intervention)
    db.commit()
    return {"message": "Success"}

@app.get("/api/stats/users")
def get_user_stats(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
        
    from sqlalchemy import func
    stats = db.query(
        models.Intervention.created_by.label("username"),
        func.count(models.Intervention.id).label("count")
    ).group_by(models.Intervention.created_by).order_by(func.count(models.Intervention.id).desc()).all()
    
    return [{"username": s.username, "count": s.count} for s in stats]

@app.post("/api/interventions/{id}/request")
def create_intervention_request(id: int, payload: schemas.RequestCreate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role == "izleyici":
        raise HTTPException(status_code=403, detail="Forbidden")
    intervention = db.query(models.Intervention).filter(models.Intervention.id == id).first()
    if not intervention:
        raise HTTPException(status_code=404, detail="Not found")
    
    # Create the request
    req = models.Request(
        intervention_id=id,
        request_type=payload.request_type,
        new_data=payload.new_data,
        status="PENDING",
        requested_by=current_user.username
    )
    db.add(req)
    
    # Mark the intervention
    intervention.flag_request = "PENDING"
    
    db.commit()
    return {"message": "Talep iletildi"}

@app.get("/api/requests")
def list_requests(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
        
    requests = db.query(models.Request, models.Intervention).join(
        models.Intervention, models.Request.intervention_id == models.Intervention.id
    ).filter(models.Request.status == "PENDING").all()
    
    # We will format this nicely for the frontend
    result = []
    for req, inv in requests:
        result.append({
            "id": req.id,
            "intervention_id": inv.id,
            "asset_id": inv.asset_id,
            "category": inv.category,
            "old_status": inv.status,
            "old_description": inv.description,
            "request_type": req.request_type,
            "new_data": req.new_data,
            "requested_by": req.requested_by,
            "created_at": req.created_at.isoformat()
        })
    return result

@app.post("/api/requests/{id}/approve")
def approve_request(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
        
    req = db.query(models.Request).filter(models.Request.id == id).first()
    if not req:
        raise HTTPException(status_code=404, detail="Not found")
        
    intervention = db.query(models.Intervention).filter(models.Intervention.id == req.intervention_id).first()
    if not intervention:
        raise HTTPException(status_code=404, detail="Intervention not found")
        
    if req.request_type == "DELETE":
        db.delete(intervention)
    elif req.request_type == "UPDATE":
        import json
        if req.new_data:
            data = json.loads(req.new_data)
            if "status" in data:
                intervention.status = data["status"]
            if "description" in data:
                intervention.description = data["description"]
            if "length_km" in data and data["length_km"]:
                intervention.length_km = float(data["length_km"])
        intervention.flag_request = None
        
    req.status = "APPROVED"
    db.commit()
    return {"message": "Success"}

@app.post("/api/requests/{id}/reject")
def reject_request(id: int, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
        
    req = db.query(models.Request).filter(models.Request.id == id).first()
    if not req:
        raise HTTPException(status_code=404, detail="Not found")
        
    intervention = db.query(models.Intervention).filter(models.Intervention.id == req.intervention_id).first()
    if intervention:
        intervention.flag_request = None
        
    req.status = "REJECTED"
    db.commit()
    return {"message": "Success"}

@app.get("/api/summary/tender-extras")
def get_tender_extras(oms: str = None, db: Session = Depends(get_db)):
    target_oms = ['HATAY METROPOL', 'KIRIKHAN', 'REYHANLI']
    if oms and oms != "Tümü":
        target_oms = [oms.upper()]
        
    excel_map = get_ihale_excel_map()
            
    query = db.query(models.Intervention, models.Line).join(
        models.Line, models.Intervention.sira_no == models.Line.sira_no
    )
    
    extra_koridor_km = 0.0
    extra_budama_adet = 0
    
    for inv, line in query.all():
        line_om = str(line.operasyon_merkezi).strip().upper()
        if line_om not in target_oms:
            continue
            
        line_hat = str(line.hat_ismi).strip().upper()
        
        koridor_excel = excel_map.get((line_om, line_hat), {}).get("koridor", "YOK")
        budama_excel = excel_map.get((line_om, line_hat), {}).get("budama", "YOK")
        
        if inv.category == "Koridor Açma":
            if koridor_excel == "YOK" or koridor_excel == "":
                extra_koridor_km += float(inv.length_km or 0)
        elif inv.category == "Ağaç Budama":
            if budama_excel == "YOK" or budama_excel == "":
                extra_budama_adet += int(inv.quantity or 1)
                
    return {
        "extra_koridor_acma_km": round(extra_koridor_km, 2),
        "extra_agac_budama_adet": extra_budama_adet
    }

@app.put("/api/users/me/dashboard-preferences")
def update_dashboard_preferences(payload: schemas.DashboardPreferencesUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    current_user.dashboard_preferences = payload.dashboard_preferences
    db.commit()
    return {"message": "Dashboard preferences updated successfully"}

@app.post("/api/summary/custom-stats")
def get_custom_stats(payload: schemas.CustomStatRequest, db: Session = Depends(get_db)):
    q = db.query(models.Line)
    if payload.type == "om":
        q = q.filter(models.Line.operasyon_merkezi == payload.name)
    else:
        q = q.filter(models.Line.ilce == payload.name)
    
    total_lines = q.count()
    if total_lines == 0:
        return []
        
    lines = q.all()
    sira_nos = [l.sira_no for l in lines]
    
    categories = ["Ağaç Budama", "Koridor Açma", "Beton Dökümü", "Güzergah Değişimi", "Operasyon Müdahalesi"]
    
    interventions = db.query(
        models.Intervention.sira_no,
        models.Intervention.category,
        models.Intervention.status
    ).filter(
        models.Intervention.sira_no.in_(sira_nos)
    ).order_by(models.Intervention.created_at.asc()).all()
    
    latest_interventions = {}
    for inv in interventions:
        key = (inv.sira_no, inv.category)
        latest_interventions[key] = inv.status
        
    results = []
    for cat in categories:
        yapildi = 0
        yapilmadi = 0
        
        for sira_no in sira_nos:
            stat = latest_interventions.get((sira_no, cat))
            if stat == "Yapıldı":
                yapildi += 1
            elif stat == "Yapılmadı":
                yapilmadi += 1
                
        results.append({
            "category": cat,
            "yapildi": yapildi,
            "yapilmadi": yapilmadi,
            "gerek_yok": total_lines - (yapildi + yapilmadi)
        })
        
    return results

@app.get("/api/users", response_model=list[schemas.UserResponse])
def get_users(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
    return db.query(models.User).all()

@app.post("/api/users", response_model=schemas.UserResponse)
def create_user(user: schemas.UserCreate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
    
    existing_user = db.query(models.User).filter(models.User.username == user.username).first()
    if existing_user:
        raise HTTPException(status_code=400, detail="Username already registered")
        
    hashed_password = auth.get_password_hash(user.password)
    db_user = models.User(
        username=user.username,
        password_hash=hashed_password,
        role=user.role,
        default_il="Tümü",
        default_oms="Tümü"
    )
    db.add(db_user)
    db.commit()
    db.refresh(db_user)
    return db_user

@app.put("/api/users/{username}", response_model=schemas.UserResponse)
def update_user(username: str, user: schemas.UserUpdate, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Forbidden")
        
    db_user = db.query(models.User).filter(models.User.username == username).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
        
    if user.password:
        db_user.password_hash = auth.get_password_hash(user.password)
    db_user.role = user.role
    
    db.commit()
    db.refresh(db_user)
    return db_user

@app.get("/api/export/preview")
def get_export_preview(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    excel_map = get_ihale_excel_map()
    toroslar_map = get_toroslar_excel_map()
    lines = db.query(models.Line).all()
    
    interventions = db.query(models.Intervention).order_by(models.Intervention.created_at.asc()).all()
    
    latest_interventions = {}
    for inv in interventions:
        key = (inv.sira_no, inv.category, inv.asset_id)
        latest_interventions[key] = inv
        
    stats_map = {}
    for key, inv in latest_interventions.items():
        sira_no = inv.sira_no
        cat = inv.category
        if sira_no not in stats_map:
            stats_map[sira_no] = {}
        if cat not in stats_map[sira_no]:
            stats_map[sira_no][cat] = {"yapildi": 0.0, "yapilacak": 0.0, "ihale": 0.0}
            
        is_yapildi = inv.status == "Yapıldı"
        # Ağaç budama vb. için adet
        qty = float(inv.quantity or 1) if cat != "Koridor Açma" else float(inv.length_km or 0)
        
        if is_yapildi:
            stats_map[sira_no][cat]["yapildi"] += qty
        else:
            stats_map[sira_no][cat]["yapilacak"] += qty

    result = []
    for line in lines:
        line_om = str(line.operasyon_merkezi).strip().upper()
        line_hat = str(line.hat_ismi).strip().upper()
        
        koridor_excel = excel_map.get((line_om, line_hat), {}).get("koridor", "YOK")
        budama_excel = excel_map.get((line_om, line_hat), {}).get("budama", "YOK")
        
        ihale_koridor = "VAR" if koridor_excel != "YOK" and koridor_excel != "" else "YOK"
        ihale_budama = "VAR" if budama_excel != "YOK" and budama_excel != "" else "YOK"
        
        sira = line.sira_no
        s_map = stats_map.get(sira, {})
        
        def get_stat(c_name, key_name):
            if key_name == "ihale":
                return "VAR" if (c_name == "Ağaç Budama" and ihale_budama == "VAR") or (c_name == "Koridor Açma" and ihale_koridor == "VAR") else "YOK"
            return s_map.get(c_name, {}).get(key_name, 0.0)

        t_data = toroslar_map.get((line_om, line_hat), {})
        
        result.append({
            "sira_no": t_data.get("sira_no", line.sira_no),
            "dagitim_sirketi": t_data.get("dagitim_sirketi", line.dagitim_sirketi),
            "il": t_data.get("il", line.il),
            "ilce": t_data.get("ilce", line.ilce),
            "operasyon_merkezi": line_om,
            "hat_ismi": line_hat,
            "gerilim_seviyesi": t_data.get("gerilim_seviyesi", line.gerilim_seviyesi),
            "hat_uzunlugu": t_data.get("hat_uzunlugu", line.hat_uzunlugu),
            "mevcut_risk": t_data.get("mevcut_risk", line.mevcut_risk),
            "planlanan_bakim": t_data.get("planlanan_bakim", ""),
            "gerceklesen_bakim": t_data.get("gerceklesen_bakim", ""),
            "siparis_no": t_data.get("siparis_no", line.siparis_no),
            "agac_budama": {
                "yapildi": get_stat("Ağaç Budama", "yapildi"),
                "yapilacak": get_stat("Ağaç Budama", "yapilacak"),
                "ihale": get_stat("Ağaç Budama", "ihale"),
            },
            "guzergah_degisimi": {
                "yapildi": get_stat("Güzergah Değişimi", "yapildi"),
                "yapilacak": get_stat("Güzergah Değişimi", "yapilacak"),
                "ihale": get_stat("Güzergah Değişimi", "ihale"),
            },
            "beton_dokumu": {
                "yapildi": get_stat("Beton Dökümü", "yapildi"),
                "yapilacak": get_stat("Beton Dökümü", "yapilacak"),
                "ihale": get_stat("Beton Dökümü", "ihale"),
            },
            "koridor_acma": {
                "yapildi": get_stat("Koridor Açma", "yapildi"),
                "yapilacak": get_stat("Koridor Açma", "yapilacak"),
                "ihale": get_stat("Koridor Açma", "ihale"),
            },
            "operasyon_mudahalesi": {
                "yapildi": get_stat("Operasyon Müdahalesi", "yapildi"),
                "yapilacak": get_stat("Operasyon Müdahalesi", "yapilacak"),
                "ihale": get_stat("Operasyon Müdahalesi", "ihale"),
            }
        })
        
    return result



def generate_excel_file(db: Session):
    excel_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'Kopya Ormanlık Alan Bakım Takip - Toroslar.xlsx')
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    lines = db.query(models.Line).all()
    line_map = {l.sira_no: l for l in lines}
    
    interventions = db.query(models.Intervention).all()
    stats = {}
    for inv in interventions:
        line_obj = line_map.get(inv.sira_no)
        if not line_obj: continue
        k1 = str(line_obj.operasyon_merkezi).strip().upper()
        k2 = str(line_obj.hat_ismi).strip().upper()
        if (k1, k2) not in stats:
            stats[(k1, k2)] = {}
        c_name = inv.category
        if c_name not in stats[(k1, k2)]:
            stats[(k1, k2)][c_name] = {"yapildi": 0.0, "yapilacak": 0.0, "ihale": 0.0}
        status = str(inv.status).lower()
        if status == "yapıldı": stats[(k1, k2)][c_name]["yapildi"] += float(inv.quantity or 1)
        elif status == "yapılmadı": stats[(k1, k2)][c_name]["yapilacak"] += float(inv.quantity or 1)
        elif status == "bekliyor": stats[(k1, k2)][c_name]["yapilacak"] += float(inv.quantity or 1)

    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    # Stil Tanımlamaları
    hdr_font = Font(bold=True)
    hdr_align = Alignment(horizontal="center", vertical="center")
    hdr_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # 2. Satır: Ana Kategori Başlıkları (Merged)
    categories = [
        ("Ağaç Budama", 13, 15),
        ("Güzergah Değişimi", 16, 18),
        ("Beton Dökümü", 19, 21),
        ("Koridor Açma", 22, 24),
        ("Operasyon Müdahalesi", 25, 27)
    ]
    
    for cat_name, start_col, end_col in categories:
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
        cell = ws.cell(row=2, column=start_col)
        cell.value = cat_name
        cell.font = hdr_font
        cell.alignment = hdr_align
        for c in range(start_col, end_col + 1):
            ws.cell(row=2, column=c).border = thin_border
            ws.cell(row=2, column=c).fill = hdr_fill

    # 3. Satır: Alt Başlıklar
    sub_headers = ["Yapılan", "Yapılacak", "İhale"] * 5
    for i, header in enumerate(sub_headers):
        col = 13 + i
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.font = hdr_font
        cell.alignment = hdr_align
        cell.border = thin_border
        if header == "Yapılan":
            cell.fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
        elif header == "Yapılacak":
            cell.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
        elif header == "İhale":
            cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for row in range(4, ws.max_row + 1):
        hat_ismi_cell = ws.cell(row=row, column=6)
        op_merkezi_cell = ws.cell(row=row, column=5)
        if not hat_ismi_cell.value or not op_merkezi_cell.value: continue
        hat = str(hat_ismi_cell.value).strip().upper()
        om = str(op_merkezi_cell.value).strip().upper()
        s_map = stats.get((om, hat), {})
        def get_stat(c_name, key_name): return s_map.get(c_name, {}).get(key_name, 0.0)
            
        ws.cell(row=row, column=13).value = get_stat("Ağaç Budama", "yapildi")
        ws.cell(row=row, column=14).value = get_stat("Ağaç Budama", "yapilacak")
        ws.cell(row=row, column=15).value = get_stat("Ağaç Budama", "ihale")
        ws.cell(row=row, column=16).value = get_stat("Güzergah Değişimi", "yapildi")
        ws.cell(row=row, column=17).value = get_stat("Güzergah Değişimi", "yapilacak")
        ws.cell(row=row, column=18).value = get_stat("Güzergah Değişimi", "ihale")
        ws.cell(row=row, column=19).value = get_stat("Beton Dökümü", "yapildi")
        ws.cell(row=row, column=20).value = get_stat("Beton Dökümü", "yapilacak")
        ws.cell(row=row, column=21).value = get_stat("Beton Dökümü", "ihale")
        ws.cell(row=row, column=22).value = get_stat("Koridor Açma", "yapildi")
        ws.cell(row=row, column=23).value = get_stat("Koridor Açma", "yapilacak")
        ws.cell(row=row, column=24).value = get_stat("Koridor Açma", "ihale")
        ws.cell(row=row, column=25).value = get_stat("Operasyon Müdahalesi", "yapildi")
        ws.cell(row=row, column=26).value = get_stat("Operasyon Müdahalesi", "yapilacak")
        ws.cell(row=row, column=27).value = get_stat("Operasyon Müdahalesi", "ihale")

        # Veri satırlarına da kenarlık ve ortalama ekleyelim ki şık dursun
        for c in range(13, 28):
            ws.cell(row=row, column=c).border = thin_border
            ws.cell(row=row, column=c).alignment = Alignment(horizontal="center")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

@app.get("/api/export/excel")
def export_excel(db: Session = Depends(get_db)):
    output = generate_excel_file(db)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Merkez_Raporu_Guncel.xlsx"}
    )



if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)

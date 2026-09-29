import openpyxl

wb = openpyxl.load_workbook('/Users/mrkantarci/Desktop/Orman Yangınları/özet rapor.xlsx', data_only=True)
ws = wb.active

non_empty_rows = 0
for row in range(4, ws.max_row + 1):
    has_data = False
    for col in range(13, 28):
        val = ws.cell(row=row, column=col).value
        if val is not None and val != 0 and val != 0.0 and str(val).strip() != '' and str(val).strip() != '0':
            has_data = True
            break
    if has_data:
        non_empty_rows += 1
        sira = ws.cell(row=row, column=1).value
        print(f"Row {row} | Sira No: {sira}")
        for col in range(13, 28):
            val = ws.cell(row=row, column=col).value
            if val is not None and val != 0 and val != 0.0 and str(val).strip() != '' and str(val).strip() != '0':
                print(f"  Col {col}: {val}")
        if non_empty_rows >= 10:
            break

print(f"Total rows with data: {non_empty_rows}")
